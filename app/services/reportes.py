"""Generacion de reportes en PDF (reportlab) y Excel (openpyxl).

Replica, por disciplina, el formato del Excel original (cantidad, P.U., total
presupuestado, peso ponderado, cantidad ejecutada, monto ejecutado, % avance),
mas una hoja de resumen general con los totales por disciplina e IVA.
Cuando `ocultar_precios` es True (rol Cliente) se omiten P.U. y montos en $.
"""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

NARANJA = colors.HexColor("#C2410C")
GRAFITO = colors.HexColor("#1F2937")
GRIS_CLARO = colors.HexColor("#F1F5F9")
VERDE = colors.HexColor("#15803D")


def _encabezados(ocultar_precios):
    base = ["#", "Descripcion", "Unidad", "Cant. Presup.", "Cant. Ejecutada", "% Avance"]
    if ocultar_precios:
        return base
    return base[:3] + ["P.U. ($)", "Total Presup. ($)", "Peso %"] + base[3:] + ["Monto Ejec. ($)"]


def _fila_partida(p, ocultar_precios):
    base = [p.numero_partida, p.descripcion, p.unidad,
            float(p.cantidad_presupuestada), float(p.cantidad_ejecutada),
            f"{float(p.porcentaje_avance) * 100:.1f}%"]
    if ocultar_precios:
        return base
    return (base[:3]
            + [float(p.precio_unitario), float(p.total_presupuestado),
               f"{float(p.peso_ponderado) * 100:.2f}%"]
            + base[3:]
            + [float(p.monto_ejecutado)])


def _excel(proyecto, ocultar_precios):
    wb = Workbook()
    encabezados = _encabezados(ocultar_precios)
    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")

    primera = True
    for disciplina in proyecto.disciplinas:
        ws = wb.active if primera else wb.create_sheet()
        primera = False
        ws.title = disciplina.nombre[:31]

        for col, titulo in enumerate(encabezados, start=1):
            celda = ws.cell(row=1, column=col, value=titulo)
            celda.fill = head_fill
            celda.font = head_font
            celda.alignment = Alignment(horizontal="center")

        for r, partida in enumerate(disciplina.partidas, start=2):
            for c, valor in enumerate(_fila_partida(partida, ocultar_precios), start=1):
                ws.cell(row=r, column=c, value=valor)

        for col in range(1, len(encabezados) + 1):
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 20

    resumen = wb.create_sheet("Resumen", 0)
    enc_resumen = ["Disciplina", "Subtotal ($)", "IVA ($)", "Total c/IVA ($)",
                   "Ejecutado ($)", "% Avance", "Peso %"] if not ocultar_precios else \
                  ["Disciplina", "% Avance", "Peso %"]
    for col, titulo in enumerate(enc_resumen, start=1):
        celda = resumen.cell(row=1, column=col, value=titulo)
        celda.fill = head_fill
        celda.font = head_font
        celda.alignment = Alignment(horizontal="center")

    total = float(proyecto.monto_total) or 1
    for r, d in enumerate(proyecto.disciplinas, start=2):
        peso = float(d.subtotal_presupuestado) / total * 100
        if ocultar_precios:
            fila = [d.nombre, f"{float(d.avance) * 100:.1f}%", f"{peso:.2f}%"]
        else:
            fila = [d.nombre, float(d.subtotal_presupuestado), float(d.iva_monto),
                    float(d.total_con_iva), float(d.monto_ejecutado),
                    f"{float(d.avance) * 100:.1f}%", f"{peso:.2f}%"]
        for c, valor in enumerate(fila, start=1):
            resumen.cell(row=r, column=c, value=valor)

    fila_total = len(proyecto.disciplinas) + 2
    if ocultar_precios:
        resumen.cell(row=fila_total, column=1, value="TOTAL GENERAL")
        resumen.cell(row=fila_total, column=2, value=f"{float(proyecto.avance_financiero) * 100:.1f}%")
    else:
        resumen.cell(row=fila_total, column=1, value="TOTAL GENERAL")
        resumen.cell(row=fila_total, column=2, value=float(proyecto.monto_total))
        resumen.cell(row=fila_total, column=4, value=float(proyecto.monto_total_con_iva))
        resumen.cell(row=fila_total, column=5, value=float(proyecto.monto_ejecutado_total))
        resumen.cell(row=fila_total, column=6, value=f"{float(proyecto.avance_financiero) * 100:.1f}%")
    for cell in resumen[fila_total]:
        cell.font = Font(bold=True)

    for col in range(1, len(enc_resumen) + 1):
        resumen.column_dimensions[resumen.cell(row=1, column=col).column_letter].width = 20

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _pdf(proyecto, ocultar_precios):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    styles["Title"].textColor = GRAFITO
    elementos = [
        Paragraph(f"Reporte de Avance de Obra &mdash; {proyecto.nombre}", styles["Title"]),
        Paragraph(f"Fecha de corte: {proyecto.fecha_corte:%d/%m/%Y} &mdash; "
                  f"Generado el {datetime.now():%d/%m/%Y %H:%M}", styles["Normal"]),
        Spacer(1, 6 * mm),
    ]

    resumen_head = ["Disciplina", "% Avance", "Peso %"] if ocultar_precios else \
                   ["Disciplina", "Subtotal ($)", "IVA ($)", "Total c/IVA ($)", "Ejecutado ($)", "% Avance"]
    total = float(proyecto.monto_total) or 1
    resumen_filas = []
    for d in proyecto.disciplinas:
        peso = float(d.subtotal_presupuestado) / total * 100
        if ocultar_precios:
            resumen_filas.append([d.nombre, f"{float(d.avance) * 100:.1f}%", f"{peso:.2f}%"])
        else:
            resumen_filas.append([
                d.nombre, f"{float(d.subtotal_presupuestado):,.2f}", f"{float(d.iva_monto):,.2f}",
                f"{float(d.total_con_iva):,.2f}", f"{float(d.monto_ejecutado):,.2f}",
                f"{float(d.avance) * 100:.1f}%",
            ])

    elementos.append(Paragraph("Resumen por disciplina", styles["Heading2"]))
    tabla_resumen = Table([resumen_head] + resumen_filas, repeatRows=1)
    tabla_resumen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), GRAFITO),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GRIS_CLARO]),
        ("GRID", (0, 0), (-1, -1), 0.4, NARANJA),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elementos.append(tabla_resumen)
    elementos.append(Spacer(1, 8 * mm))

    for d in proyecto.disciplinas:
        elementos.append(Paragraph(d.nombre, styles["Heading2"]))
        encabezados = _encabezados(ocultar_precios)
        filas = [_fila_partida(p, ocultar_precios) for p in d.partidas]
        data = [encabezados] + [[str(x) for x in f] for f in filas]
        tabla = Table(data, repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NARANJA),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, GRIS_CLARO]),
            ("GRID", (0, 0), (-1, -1), 0.4, GRAFITO),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elementos.append(tabla)
        elementos.append(Spacer(1, 6 * mm))

    doc.build(elementos)
    buffer.seek(0)
    return buffer


def reporte_avance(proyecto, formato="pdf", ocultar_precios=False):
    if formato == "excel":
        return _excel(proyecto, ocultar_precios), "xlsx"
    return _pdf(proyecto, ocultar_precios), "pdf"


def _pdf_partidas_disciplina(disciplina):
    """Ficha PDF de una disciplina: partidas + bloque de Valuacion (seccion 3.1)."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    styles["Title"].textColor = GRAFITO

    encabezados = ["#", "Descripcion", "Unidad", "Cant.\nPresup.", "P.U. ($)",
                   "Total\nPresup. ($)", "% Avance", "Valuacion\nTotal ($)"]
    filas = []
    for p in disciplina.partidas:
        filas.append([
            str(p.numero_partida), p.descripcion, p.unidad,
            f"{float(p.cantidad_presupuestada):,.2f}", f"{float(p.precio_unitario):,.2f}",
            f"{float(p.total_presupuestado):,.2f}", f"{float(p.porcentaje_avance) * 100:.1f}%",
            f"{float(p.valuacion_total):,.2f}",
        ])
    filas.append([
        "", "TOTAL", "", "", "",
        f"{float(disciplina.subtotal_presupuestado):,.2f}", f"{float(disciplina.avance) * 100:.1f}%",
        f"{float(disciplina.valuacion_total):,.2f}",
    ])

    tabla = Table([encabezados] + filas, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NARANJA),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, GRIS_CLARO]),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
        ("GRID", (0, 0), (-1, -1), 0.4, GRAFITO),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))

    elementos = [
        Paragraph(f"Ficha de partidas &mdash; {disciplina.nombre}", styles["Title"]),
        Paragraph(f"{disciplina.proyecto.nombre} &mdash; Generado el {datetime.now():%d/%m/%Y %H:%M}",
                   styles["Normal"]),
        Spacer(1, 6 * mm),
        tabla,
    ]
    doc.build(elementos)
    buffer.seek(0)
    return buffer


def reporte_partidas_disciplina(disciplina):
    return _pdf_partidas_disciplina(disciplina), "pdf"
